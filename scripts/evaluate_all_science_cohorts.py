"""Compare pooled science curves without publishing student-level data."""
import argparse
import json
from pathlib import Path

from openpyxl import load_workbook

from evaluate_science_cohorts import fit, measure
from validate_science_cohorts import extract, numeric

ROOT = Path(__file__).resolve().parent.parent
SHEETS = {'Biology': ('Biology', 18, 19), 'Chemistry': ('Chem', 22, 23), 'Physics': ('Physics', 18, 19)}


def original_pairs(workbook, subject):
    sheet, mock_col, actual_col = SHEETS[subject]
    pairs = []
    for row in workbook[sheet].iter_rows(min_row=2, max_row=250, values_only=True):
        mock, actual = row[mock_col - 1], row[actual_col - 1]
        if numeric(mock) and numeric(actual):
            if not (0 <= mock <= 50 and 0 <= actual <= 50):
                raise ValueError(f'{sheet}: result outside 0–50')
            pairs.append((mock * 2, actual * 2))
    return pairs


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('original')
    parser.add_argument('results_2024')
    parser.add_argument('results_2025')
    args = parser.parse_args()
    old = json.loads((ROOT / 'dist/mock-models.json').read_text())['subjects']
    workbook = load_workbook(args.original, read_only=True, data_only=True)
    result = {'description': 'Aggregate-only pooled science comparison. Earlier workbook is mostly completion year 2022 despite its 2023 filename. The pooled fit uses all three cohorts; its errors on those cohorts are in-sample.', 'subjects': {}}
    try:
        for subject in SHEETS:
            earlier = original_pairs(workbook, subject)
            if len(earlier) != old[subject]['n']:
                raise ValueError(f'{subject}: original pair count differs from current model')
            year24 = extract(args.results_2024, 2024, subject)
            year25 = extract(args.results_2025, 2025, subject)
            pooled = fit(earlier + year24 + year25)
            train_early24 = fit(earlier + year24)
            train_early25 = fit(earlier + year25)
            result['subjects'][subject] = {
                'counts': {'earlier': len(earlier), '2024': len(year24), '2025': len(year25)},
                'pooledParameters': pooled,
                'pooledInSample': {label: measure(pairs, pooled) for label, pairs in [('earlier', earlier), ('2024', year24), ('2025', year25)]},
                'currentHeldOut': {label: measure(pairs, old[subject]['parameters']) for label, pairs in [('2024', year24), ('2025', year25)]},
                'earlierPlus2024Test2025': measure(year25, train_early24),
                'earlierPlus2025Test2024': measure(year24, train_early25),
            }
    finally:
        workbook.close()
    for path in (ROOT / 'analysis/science-all-cohorts.json', ROOT / 'dist/science-all-cohorts.json'):
        path.write_text(json.dumps(result, indent=2, ensure_ascii=False) + '\n')
    for subject, item in result['subjects'].items():
        print(subject, 'pooled', sum(item['counts'].values()), '2025 held-out RMSE', item['earlierPlus2024Test2025']['rmse'])


if __name__ == '__main__':
    main()
