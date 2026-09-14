"""Add aggregate-only 2024/2025 science validation to existing mock models.

Student names, codes and score pairs are read locally and never written out.
Run with the 2024 and 2025 results workbooks as arguments.
"""
import argparse
import json
import math
from pathlib import Path

from openpyxl import load_workbook


def numeric(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def extract(path, year, subject):
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[subject]
    first = 4 if year == 2024 or subject != 'Chemistry' else 2
    pairs = []
    seen = set()
    for row in sheet.iter_rows(min_row=first, values_only=True):
        if year == 2024:
            mock, actual = row[6:8]
        elif subject == 'Biology':
            # The displayed Mock/50 is rounded; source total /84 retains precision.
            mock = row[4] / 84 * 50 if len(row) > 4 and numeric(row[4]) else None
            actual = row[8] if len(row) > 8 else None
        elif subject == 'Chemistry':
            mock, actual = row[7:9]
        else:
            mock, actual = row[10:12]
        if not (numeric(mock) and numeric(actual)):
            continue
        if not (0 <= mock <= 50 and 0 <= actual <= 50):
            raise ValueError(f'{year} {subject}: result outside 0–50')
        identifier = str(row[0]).strip()
        if identifier in seen:
            raise ValueError(f'{year} {subject}: duplicate student identifier')
        seen.add(identifier)
        pairs.append((mock * 2, actual * 2))
    workbook.close()
    if len(pairs) < 10:
        raise ValueError(f'{year} {subject}: unexpectedly few paired results')
    return pairs


def metrics(pairs, parameters):
    L, A, k = parameters
    errors = [L / (1 + A * math.exp(-k * mock)) - actual for mock, actual in pairs]
    return {
        'n': len(errors),
        'mae': round(sum(map(abs, errors)) / len(errors), 2),
        'rmse': round(math.sqrt(sum(error * error for error in errors) / len(errors)), 2),
        'bias': round(sum(errors) / len(errors), 2),
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('results_2024')
    parser.add_argument('results_2025')
    args = parser.parse_args()
    path = Path(__file__).resolve().parent.parent / 'dist/mock-models.json'
    catalog = json.loads(path.read_text())
    for subject in ('Biology', 'Chemistry', 'Physics'):
        model = catalog['subjects'][subject]
        model['independentValidation'] = {
            str(year): metrics(extract(getattr(args, f'results_{year}'), year, subject), model['parameters'])
            for year in (2024, 2025)
        }
        combined_pairs = [pair for year in (2024, 2025)
                          for pair in extract(getattr(args, f'results_{year}'), year, subject)]
        L, A, k = model['parameters']
        absolute_errors = sorted(abs(L / (1 + A * math.exp(-k * mock)) - actual)
                                 for mock, actual in combined_pairs)
        model['historicalError80'] = {
            'n': len(combined_pairs),
            'absoluteErrorPP': round(absolute_errors[math.ceil(.8 * len(absolute_errors)) - 1], 2),
            'mockRange': [round(min(mock for mock, _ in combined_pairs), 2),
                          round(max(mock for mock, _ in combined_pairs), 2)],
            'cohorts': ['2024', '2025'],
        }
        model['validationNote'] = 'Independent school cohorts; metrics use unrounded predicted percentages and actual external percentages. These records were not used to fit this curve.'
        if subject == 'Chemistry':
            warning = 'Independent checks show a cohort shift: the model averaged 5.28 percentage points low in 2024 and 4.18 points high in 2025.'
            if warning not in model['warnings']:
                model['warnings'].append(warning)
        print(subject, model['independentValidation'])
    catalog['validation'] = 'Original leave-one-out training results plus independent 2024 and 2025 science-cohort checks. Errors are percentage points, not prediction intervals.'
    note = ' Biology, Chemistry and Physics were independently checked on the supplied 2024 and 2025 science results. The ZIP and separate 2025 mock workbook overlap those workbooks and were not counted again.'
    if note not in catalog['sourceNotes']:
        catalog['sourceNotes'] += note
    path.write_text(json.dumps(catalog, indent=2, ensure_ascii=False, allow_nan=False) + '\n')


if __name__ == '__main__':
    main()
