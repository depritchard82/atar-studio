"""Fit aggregate-only Mathematics mock converters from the supplied combined workbook."""
import argparse
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from evaluate_science_cohorts import fit, predict

ROOT = Path(__file__).resolve().parent.parent
SHEETS = {
    'Mathematical Methods': 'MM Combined',
    'General Mathematics': 'GM Combined',
    'Specialist Mathematics': 'SM Combined',
}


def pairs(sheet):
    values = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        mock, actual = row[:2]
        if not (isinstance(mock, (int, float)) and isinstance(actual, (int, float))):
            continue
        if not (0 <= mock <= 1 and 0 <= actual <= 1):
            raise ValueError(f'{sheet.title}: expected decimal percentages from 0 to 1')
        values.append((mock * 100, actual * 100))
    if len(values) < 20:
        raise ValueError(f'{sheet.title}: unexpectedly few paired results')
    return values


def build(subject, values, source, sheet):
    data = np.asarray(values, dtype=float)
    x, y = data.T
    parameters = fit(values)
    residual = predict(x, parameters) - y
    cv_errors, baseline_errors = [], []
    for i in range(len(values)):
        keep = np.arange(len(values)) != i
        cv_errors.append(float(predict(np.asarray([x[i]]), fit(data[keep].tolist()))[0] - y[i]))
        baseline_errors.append(float(y[keep].mean() - y[i]))
    return {
        'subject': subject,
        'n': len(values),
        'externalMaximum': 50,
        'mockRange': [float(x.min()), float(x.max())],
        'source': source,
        'sheet': sheet,
        'mapping': 'A mock percentage (EM), B real external percentage (EA); decimal percentages converted to 0–100',
        'cohort': 'Combined source; component years not supplied',
        'enabled': True,
        'status': 'Provisional',
        'warnings': ['The combined workbook does not label component years, so future-year performance cannot yet be measured.',
                     'A 100% endpoint constraint was considered but not used; it did not consistently improve validation and required implausibly high asymptotes for some subjects.'],
        'parameters': parameters,
        'cvMAE': float(np.mean(np.abs(cv_errors))),
        'cvRMSE': float(np.sqrt(np.mean(np.square(cv_errors)))),
        'baselineRMSE': float(np.sqrt(np.mean(np.square(baseline_errors)))),
        'r2': float(1 - np.sum(residual ** 2) / np.sum((y - y.mean()) ** 2)),
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('workbook')
    args = parser.parse_args()
    source = Path(args.workbook).name
    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    catalog_path = ROOT / 'dist/mock-models.json'
    catalog = json.loads(catalog_path.read_text())
    summary = {'description': 'Aggregate-only Mathematics model results; no student records or score pairs', 'subjects': {}}
    try:
        for subject, sheet in SHEETS.items():
            model = build(subject, pairs(workbook[sheet]), source, sheet)
            catalog['subjects'][subject] = model
            summary['subjects'][subject] = {key: model[key] for key in ('n', 'mockRange', 'parameters', 'cvMAE', 'cvRMSE', 'baselineRMSE', 'r2')}
    finally:
        workbook.close()
    catalog['validation'] += ' Mathematics curves use leave-one-student-out refitting on the supplied combined workbook; component cohort years were not supplied.'
    catalog['sourceNotes'] += ' Mathematics uses the supplied combined workbook as the definitive source; earlier Mathematics rows are not added again because the combined workbook does not expose identifiers or component years for reliable deduplication.'
    catalog_path.write_text(json.dumps(catalog, indent=2, ensure_ascii=False, allow_nan=False) + '\n')
    output = ROOT / 'analysis/mathematics-combined-models.json'
    output.parent.mkdir(exist_ok=True)
    output.write_text(json.dumps(summary, indent=2, ensure_ascii=False, allow_nan=False) + '\n')
    for subject, item in summary['subjects'].items():
        print(subject, item['n'], 'LOO RMSE', round(item['cvRMSE'], 2))


if __name__ == '__main__':
    main()
